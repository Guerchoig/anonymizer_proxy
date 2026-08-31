from docx import Document

doc = Document(r'c:\Test\anonymizer_proxy\docs\KP_IRIS.result.docx')

print('=== АНАЛИЗ ДОКУМЕНТА ===')
print()

# Считаем таблицы и их размеры
print(f'Общее количество таблиц: {len(doc.tables)}')
for i, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    print(f'  Таблица {i}: {rows} строк x {cols} столбцов')

print()

# Извлекаем ключевую информацию из таблиц
summary = []
for i, table in enumerate(doc.tables):
    if len(table.rows) > 1:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        summary.append({'table_num': i + 1, 'header': header})

print('Найдено многомерных таблиц:')
for item in summary:
    print(f'  - Таблица {item["table_num"]}: {len(item["header"])} столбцов')
    for h in item['header'][:3]:
        print(f'      {h}')

print()
print('=' * 60)
print('КРАТКОЕ СОДЕРЖАНИЕ ДОКУМЕНТА:')
print('=' * 60)
print()

# Проверяем наличие текстовых блоков
has_text = False
for para in doc.paragraphs[:100]:
    if len(para.text.strip()) > 20:
        has_text = True
        text_snippet = para.text.strip()[:200]
        if len(para.text.strip()) > 200:
            text_snippet += '...'
        print(f'  • {text_snippet}')

if not has_text:
    print('Текстовое содержание ограничено таблицами.')

print()
print('=' * 60)
print('ДОПОЛНИТЕЛЬНАЯ ИНФОРМАЦИЯ:')
print('=' * 60)
print(f'  • Всего абзацев: {len(doc.paragraphs)}')
print(f'  • Всего таблиц: {len(doc.tables)}')
