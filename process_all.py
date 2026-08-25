from docx import Document
from add_col import add_column_to_table
from add_template import add_template_appendix


def process_kp_iris():
    doc = Document('c:/Test/anonymizer_proxy/docs/KP_IRIS.anonymized.docx')
    print("Обработка KP_IRIS...")
    for table in doc.tables:
        add_column_to_table(table)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Приложения:":
            for j in range(i+1, min(i+10, len(doc.paragraphs))):
                if "Приложение 3" in doc.paragraphs[j].text:
                    new_p = doc.add_paragraph(
                        "Приложение 4. Шаблон структуры коммерческого предложения."
                    )
                    doc.paragraphs[j]._element.addnext(new_p._element)
                    break
            break
    add_template_appendix(doc, 4)
    doc.save('c:/Test/anonymizer_proxy/docs/KP_IRIS.result.docx')
    print("  OK: KP_IRIS.result.docx")


def process_kp_do_01():
    doc = Document('c:/Test/anonymizer_proxy/docs/КП ДО 01.anonymized.docx')
    print("Обработка КП ДО 01...")
    for table in doc.tables:
        add_column_to_table(table)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("ИНФОРМАЦИЯ О КОМПАНИИ") and "\t" in p.text:
            new_p = doc.add_paragraph("ПРИЛОЖЕНИЯ")
            doc.paragraphs[i]._element.addnext(new_p._element)
            break
    add_template_appendix(doc, 1)
    doc.save('c:/Test/anonymizer_proxy/docs/КП ДО 01.result.docx')
    print("  OK: КП ДО 01.result.docx")


if __name__ == '__main__':
    process_kp_iris()
    process_kp_do_01()
    print("Готово!")
