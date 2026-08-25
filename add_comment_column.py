from docx import Document
from docx.oxml import OxmlElement
from copy import deepcopy


def add_column_to_table(table, header_text="Комментарий"):
    """Добавляет столбец справа в таблицу с указанным заголовком.
    Работает напрямую с XML, корректно обрабатывая merged cells."""
    
    tbl = table._tbl
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    
    # Обновляем tblGrid — добавляем новый gridCol
    tbl_grid = tbl.find(f'{ns}tblGrid')
    if tbl_grid is not None:
        grid_cols = tbl_grid.findall(f'{ns}gridCol')
        if grid_cols:
            # Копируем ширину последнего столбца
            new_grid_col = deepcopy(grid_cols[-1])
            tbl_grid.append(new_grid_col)
        else:
            new_grid_col = OxmlElement('w:gridCol')
            tbl_grid.append(new_grid_col)
    
    rows = tbl.findall(f'{ns}tr')
    
    for row_idx, tr in enumerate(rows):
        # Находим все tc в строке
        tcs = tr.findall(f'{ns}tc')
        
        if tcs:
            # Берём последнюю ячейку как образец для копирования форматирования
            last_tc = tcs[-1]
            new_tc = deepcopy(last_tc)
            
            # Убираем gridSpan если есть (новая ячейка — отдельная)
            tc_pr = new_tc.find(f'{ns}tcPr')
            if tc_pr is not None:
                grid_span = tc_pr.find(f'{ns}gridSpan')
                if grid_span is not None:
                    tc_pr.remove(grid_span)
                # Убираем vMerge если есть
                v_merge = tc_pr.find(f'{ns}vMerge')
                if v_merge is not None:
                    tc_pr.remove(v_merge)
            
            # Удаляем все параграфы из новой ячейки
            for p in list(new_tc.findall(f'{ns}p')):
                new_tc.remove(p)
        else:
            new_tc = OxmlElement('w:tc')
        
        # Добавляем пустой параграф
        new_p = OxmlElement('w:p')
        new_tc.append(new_p)
        
        # Для первой строки добавляем заголовок
        if row_idx == 0:
            new_r = OxmlElement('w:r')
            new_t = OxmlElement('w:t')
            new_t.text = header_text
            new_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            new_r.append(new_t)
            new_p.append(new_r)
        
        # Добавляем новую ячейку в конец строки
        tr.append(new_tc)


def process_document(input_path, output_path):
    """Обрабатывает документ: добавляет столбец в каждую таблицу"""
    doc = Document(input_path)
    
    print(f"Обработка: {input_path}")
    print(f"Найдено таблиц: {len(doc.tables)}")
    
    for idx, table in enumerate(doc.tables):
        print(f"  Таблица {idx}: {len(table.rows)} строк x {len(table.columns)} столбцов")
        add_column_to_table(table)
        print(f"    -> Добавлен столбец, теперь: {len(table.rows)} строк x {len(table.columns)} столбцов")
    
    doc.save(output_path)
    print(f"Сохранено: {output_path}\n")


# Обрабатываем оба файла
files = [
    ('c:/Test/anonymizer_proxy/docs/KP_IRIS.anonymized.docx', 
     'c:/Test/anonymizer_proxy/docs/KP_IRIS.result.docx'),
    ('c:/Test/anonymizer_proxy/docs/КП ДО 01.anonymized.docx', 
     'c:/Test/anonymizer_proxy/docs/КП ДО 01.result.docx')
]

for input_path, output_path in files:
    process_document(input_path, output_path)

print("Готово!")
