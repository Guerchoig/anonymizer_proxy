#!/usr/bin/env python
# -*- coding: utf-8 -*-
import docx
import difflib

def extract_content(path):
    """Извлечение содержимого из docx файла"""
    d = docx.Document(path)
    
    # Извлечение параграфов
    paragraphs = []
    for i, p in enumerate(d.paragraphs):
        paragraphs.append(f"{i}: [{p.style.name}] {p.text}")
    
    # Извлечение таблиц
    tables = []
    for ti, table in enumerate(d.tables):
        rows_data = []
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            rows_data.append(cells)
        tables.append((ti, rows_data))
    
    return paragraphs, tables

def compare_paragraphs(p1, p2):
    """Сравнение параграфов и возврат различий"""
    differences = []
    matcher = difflib.SequenceMatcher(None, p1, p2)
    
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            continue
        elif tag == 'replace':
            for idx in range(max(i2-i1, j2-j1)):
                old = p1[i1+idx] if i1+idx < i2 else "<отсутствует>"
                new = p2[j1+idx] if j1+idx < j2 else "<отсутствует>"
                if old != new:
                    differences.append({'type': 'replace', 'old': old, 'new': new})
        elif tag == 'delete':
            for idx in range(i1, i2):
                differences.append({'type': 'delete', 'old': p1[idx], 'new': None})
        elif tag == 'insert':
            for idx in range(j1, j2):
                differences.append({'type': 'insert', 'old': None, 'new': p2[idx]})
    
    return differences

def compare_tables(t1, t2):
    """Сравнение таблиц и возврат различий"""
    differences = []
    if len(t1) != len(t2):
        differences.append({'type': 'table_count', 'old': len(t1), 'new': len(t2)})
        return differences
    
    for (ti1, rows1), (ti2, rows2) in zip(t1, t2):
        if len(rows1) != len(rows2):
            differences.append({'type': 'table_rows', 'table_index': ti1, 'old': len(rows1), 'new': len(rows2)})
            continue
        for ri, (row1, row2) in enumerate(zip(rows1, rows2)):
            if row1 != row2:
                differences.append({'type': 'table_cell', 'table_index': ti1, 'row_index': ri, 'old': row1, 'new': row2})
    
    return differences

def generate_report(file1, file2, output_path):
    """Генерация отчёта в формате Markdown"""
    p1, t1 = extract_content(file1)
    p2, t2 = extract_content(file2)
    
    para_diffs = compare_paragraphs(p1, p2)
    table_diffs = compare_tables(t1, t2)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("# Отчёт о сравнении документов\n\n")
        f.write(f"**Файл 1:** `{file1}`\n")
        f.write(f"**Файл 2:** `{file2}`\n\n")
        
        f.write("## Общая информация\n\n")
        f.write(f"- Количество параграфов в файле 1: {len(p1)}\n")
        f.write(f"- Количество параграфов в файле 2: {len(p2)}\n")
        f.write(f"- Количество таблиц в файле 1: {len(t1)}\n")
        f.write(f"- Количество таблиц в файле 2: {len(t2)}\n\n")
        
        f.write("## Различия в параграфах\n\n")
        if not para_diffs:
            f.write("*Различий не обнаружено*\n\n")
        else:
            f.write(f"Найдено различий: {len(para_diffs)}\n\n")
            for i, diff in enumerate(para_diffs, 1):
                f.write(f"### Различие #{i}\n\n")
                if diff['type'] == 'replace':
                    f.write("**Тип:** Замена\n\n")
                    f.write(f"**Было:**\n```\n{diff['old']}\n```\n\n")
                    f.write(f"**Стало:**\n```\n{diff['new']}\n```\n\n")
                elif diff['type'] == 'delete':
                    f.write("**Тип:** Удаление\n\n")
                    f.write(f"**Было:**\n```\n{diff['old']}\n```\n\n")
                elif diff['type'] == 'insert':
                    f.write("**Тип:** Добавление\n\n")
                    f.write(f"**Стало:**\n```\n{diff['new']}\n```\n\n")
                f.write("---\n\n")
        
        f.write("## Различия в таблицах\n\n")
        if not table_diffs:
            f.write("*Различий не обнаружено*\n\n")
        else:
            f.write(f"Найдено различий: {len(table_diffs)}\n\n")
            for i, diff in enumerate(table_diffs, 1):
                f.write(f"### Различие #{i}\n\n")
                if diff['type'] == 'table_count':
                    f.write("**Тип:** Изменение количества таблиц\n\n")
                    f.write(f"**Было:** {diff['old']} таблиц\n\n")
                    f.write(f"**Стало:** {diff['new']} таблиц\n\n")
                elif diff['type'] == 'table_rows':
                    f.write(f"**Тип:** Изменение количества строк в таблице #{diff['table_index']}\n\n")
                    f.write(f"**Было:** {diff['old']} строк\n\n")
                    f.write(f"**Стало:** {diff['new']} строк\n\n")
                elif diff['type'] == 'table_cell':
                    f.write(f"**Тип:** Изменение ячеек в таблице #{diff['table_index']}, строка #{diff['row_index']}\n\n")
                    f.write(f"**Было:**\n```\n{diff['old']}\n```\n\n")
                    f.write(f"**Стало:**\n```\n{diff['new']}\n```\n\n")
                f.write("---\n\n")
        
        f.write("## Заключение\n\n")
        total_diffs = len(para_diffs) + len(table_diffs)
        if total_diffs == 0:
            f.write("Документы идентичны.\n")
        else:
            f.write(f"Всего обнаружено {total_diffs} различий.\n")
    
    print(f"Отчёт создан: {output_path}")
    print(f"Различий в параграфах: {len(para_diffs)}")
    print(f"Различий в таблицах: {len(table_diffs)}")

if __name__ == '__main__':
    file1 = r'c:\Test\anonymizer_proxy\docs\КП ДО 01.docx'
    file2 = r'c:\Test\anonymizer_proxy\docs\КП ДО 01.result.docx'
    output = r'c:\Test\anonymizer_proxy\docs\comparison_report.md'
    generate_report(file1, file2, output)
