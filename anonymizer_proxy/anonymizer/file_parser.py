"""
Парсеры файлов для извлечения текста
Поддерживаемые форматы: DOCX, XLSX, XML (MS Project)
"""
import io
import os
import re
import zipfile
import posixpath
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Optional
from dataclasses import dataclass, field

from lxml import etree as _lxml

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ..config import logger


@dataclass
class ParsedContent:
    """Результат парсинга файла"""
    text: str = ""
    structure: dict = field(default_factory=dict)  # Сохранение структуры для обратной сборки
    metadata: dict = field(default_factory=dict)
    markdown: str = ""  # Форматированное markdown-представление (с таблицами)


def _normalize_line(text: str) -> str:
    """Одна строка текста: без переносов (иначе обратная сборка по
    split('\n') даст смещение)"""
    return re.sub(r"\s*\n+\s*", " ", text or "").strip()


def _cell_to_line(value) -> str:
    """Значение ячейки Excel -> одна строка текста (пустая для пустых)"""
    if value is None:
        return ""
    return _normalize_line(str(value))


def _restore_cell_value(original, part: str):
    """Значение для записи обратно в ячейку XLSX при сборке.

    Если исходная ячейка была ЧИСЛОВОЙ (int/float, не bool/дата), а новый
    текст парсится как число — записываем число, а не строку: иначе после
    де-анонимизации суммы превращаются в текст («386887.422» строкой,
    точка вместо локальной запятой, сортировка/формулы ломаются;
    багрепорт 2026-09-10). Допускаются: разрядные пробелы, запятая как
    десятичный разделитель и символы валют (₽/$/€/«руб.» — в том числе
    от рендера «как в Excel», см. _render_numeric_cell).
    Во всех остальных случаях возвращаем текст как есть."""
    if isinstance(original, (int, float)) and not isinstance(original, bool):
        cleaned = part.strip()
        for ch in ("\u00A0", " ", "\u2009"):
            cleaned = cleaned.replace(ch, "")
        cleaned = re.sub(r"(?i)[₽$€¥£]|руб\w*|р\.", "", cleaned)
        if cleaned:
            normalized = cleaned.replace(",", ".")
            try:
                num = float(normalized)
            except ValueError:
                return part
            if num.is_integer() and "." not in cleaned and "," not in cleaned:
                return int(num)
            return num
    return part


def _fmt_cell_value(value) -> str:
    """Значение ячейки Excel в читаемом виде (markdown, read-column).

    Даты — ISO (YYYY-MM-DD, без шума «00:00:00»), escape _x000D_ (CR,
    попадающий в текст при Excel-экспорте многострочных заголовков) —
    убирается, переносы строк сворачиваются (markdown-таблица обязана
    быть однострочной). Используется ТОЛЬКО для отображения: сегменты
    parse/assemble (_cell_to_line) остаются как есть — обратная сборка
    по 1:1-сравнению не ломается.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return _normalize_line(re.sub(r"_x000D_", "", str(value)))


# Символы валют в кодах форматов Excel: [$₽-419], \$, _₽, литералы
# «руб»/«р.»; учётные (Accounting/«Финансовый») и денежные (Currency)
# форматы — именно они обозначают суммы (см. справку Microsoft:
# Accounting/Currency formats; коды вида _-* #,##0.00\ _₽_-)
_CURRENCY_IN_FORMAT_RE = re.compile(
    r"\[\$([^\]-]+)[^\]]*\]"       # [$€-407] / [$₽-419]
    r"|\\([₽$€¥£])"                # \₽ (экранированный символ)
    r"|(?<![0-9#])(₽|\$|€|¥|£)"    # одиночный символ
    r"|\"((?:руб|р)\.?|руб\w*|RUB)\"",  # текстовый литерал «руб.»/"руб"
)


def _currency_from_format(fmt: str) -> str:
    """Извлечь символ валюты из кода формата ячейки ('' — не денежный)."""
    if not fmt:
        return ""
    m = _CURRENCY_IN_FORMAT_RE.search(fmt)
    if not m:
        return ""
    return next(g for g in m.groups() if g)


def _group_digits(int_part: str) -> str:
    """Разрядная группировка целой части пробелами: 1586238 → 1 586 238."""
    out = []
    while len(int_part) > 3:
        out.append(int_part[-3:])
        int_part = int_part[:-3]
    out.append(int_part)
    return " ".join(reversed(out))


def _render_numeric_cell(value, number_format: str) -> Optional[str]:
    """Текстовое представление числовой ячейки «как в Excel».

    Для ячеек с денежным (Accounting/«Финансовый», Currency) или
    группированным форматом значение рендерится с разрядными пробелами,
    запятой-разделителем и символом валюты: str(float) («1586238.422»)
    не матчится ни одним MONEY-паттерном (нет ни группировки, ни валюты),
    и такие суммы не анонимизируются (багрепорт 2026-09-11, формат
    «Финансовый»: из 9 сумм скрылись только 2, пойманные GLiNER).
    Returns None, если рендер к деньгам/группировке не применим."""
    if isinstance(value, bool) or isinstance(value, (datetime, date)):
        return None
    if not isinstance(value, (int, float)):
        return None
    fmt = (number_format or "").lower()
    has_currency = bool(_currency_from_format(fmt))
    has_grouping = "#,##" in fmt or "# ##" in fmt
    if not (has_currency or has_grouping):
        return None
    negative = value < 0
    s = repr(abs(value)) if isinstance(value, float) else str(abs(value))
    if "e" in s or "E" in s:  # экспоненциальная запись — не рендерим
        return None
    int_part, _, frac_part = s.partition(".")
    grouped = _group_digits(int_part)
    text = f"{grouped},{frac_part}" if frac_part else grouped
    if negative:
        text = "-" + text
    currency = _currency_from_format(fmt)
    if currency:
        text = f"{text} {currency}"
    return text


def _xlsx_cell_text(cell) -> str:
    """Текст сегмента XLSX-ячейки: числа с денежным/группированным
    форматом — «как в Excel» (см. _render_numeric_cell), остальное —
    как раньше (_cell_to_line). Используется и в parse, и в assemble —
    сравнение 1:1 остаётся консистентным."""
    rendered = _render_numeric_cell(cell.value, cell.number_format or "")
    if rendered is not None:
        return rendered
    return _cell_to_line(cell.value)


def _rows_to_markdown(rows: list[list[str]]) -> str:
    """Список списков строк -> markdown-таблица (первая строка — заголовок)."""
    if not rows:
        return ""
    header = rows[0]
    col_count = max((len(r) for r in rows), default=len(header))
    lines = ["| " + " | ".join((header + [""] * col_count)[:col_count]) + " |"]
    lines.append("|" + " --- |" * col_count)
    for r in rows[1:]:
        cells = (r + [""] * col_count)[:col_count]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _nearest_txbx(node):
    """Ближайший контейнер w:txbxContent для узла (или None)."""
    for anc in node.iterancestors(qn("w:txbxContent")):
        return anc
    return None


def _scrub_docx_metadata(doc) -> None:
    """Обезличить метаданные пакета DOCX.

    docProps/core.xml хранит автора и даты правок, docProps/app.xml —
    компанию, менеджера и гиперссылки заголовков (HLinks): в
    анонимизированной копии это такая же утечка PII, как текст документа.
    Вызывается только при АНОНИМИЗИРУЮЩЕЙ сборке (не при де-анонимизации).
    """
    try:
        cp = doc.core_properties
        cp.author = ""
        cp.last_modified_by = ""
        cp.comments = ""
        cp.subject = ""
        cp.keywords = ""
        cp.category = ""
    except Exception as exc:  # noqa: BLE001 - метаданные не должны ронять сборку
        logger.warning("Не удалось очистить core-properties: %s", exc)

    try:
        for part in doc.part.package.iter_parts():
            if str(part.partname) != "/docProps/app.xml":
                continue
            xml = part.blob.decode("utf-8", errors="ignore")
            cleaned = re.sub(r"<HLinks>.*?</HLinks>", "", xml, flags=re.S)
            cleaned = re.sub(r"<Company>.*?</Company>", "", cleaned, flags=re.S)
            cleaned = re.sub(r"<Manager>.*?</Manager>", "", cleaned, flags=re.S)
            if cleaned != xml:
                part._blob = cleaned.encode("utf-8")
    except Exception as exc:  # noqa: BLE001
        logger.warning("Не удалось очистить docProps/app.xml: %s", exc)


def _p_element_text(p_el) -> str:
    """Полный текст абзаца по его XML-элементу: включает содержимое
    гиперссылок (w:hyperlink), которое теряет python-docx .text, и
    исключает текст чужих/вложенных фигур — у абзаца внутри фигуры
    учитывается только её собственный w:txbxContent."""
    own_box = _nearest_txbx(p_el)
    parts = []
    for node in p_el.iter(qn("w:t")):
        if _nearest_txbx(node) is not own_box:
            continue
        parts.append(node.text or "")
    return "".join(parts)


def _table_to_markdown(table) -> str:
    """docx Table -> markdown-таблица (сохраняет сетку ячеек)."""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = "\n".join(
                _p_element_text(p._p) for p in cell.paragraphs)
            cells.append(_normalize_line(cell_text))
        rows.append(cells)
    return _rows_to_markdown(rows)


def _table_to_markdown_deep(table) -> str:
    """Markdown таблицы вместе с вложенными в её ячейки таблицами."""
    parts = [_table_to_markdown(table)]
    for row in table.rows:
        for cell in row.cells:
            for ntab in _cell_tables(cell):
                parts.append(_table_to_markdown_deep(ntab))
    return "\n\n".join(parts)


def _cell_tables(cell) -> list:
    """Вложенные таблицы ячейки (python-docx 1.2+: _Cell.tables).

    Вложенные таблицы не видны ни doc.tables (только верхний уровень), ни
    cell.paragraphs — титульные листы («УТВЕРЖДАЮ», шапки) оставались
    неанонимизированными (багрепорт 2026-09-08).
    """
    try:
        return list(cell.tables)
    except Exception:  # noqa: BLE001 - битая вложенная таблица не роняет парсинг
        return []


def _emit_nested_cell_paras(cell, path, kind, part,
                            segments, structure, paras_out=None) -> None:
    """Рекурсивно добавить сегменты абзацев вложенных таблиц ячейки.

    path — цепочка координат [[t, r, c], ...]: первый триплет — таблица
    верхнего уровня (doc.tables / hf.tables) + строка + ячейка, каждый
    следующий — вложенная таблица внутри текущей ячейки + строка + ячейка.
    paras_out — куда дополнительно сложить тексты (entry["paras"] блока
    markdown).
    """
    for nt_idx, ntab in enumerate(_cell_tables(cell)):
        for r_idx, row in enumerate(ntab.rows):
            for c_idx, ncell in enumerate(row.cells):
                for p_idx, para in enumerate(ncell.paragraphs):
                    text = _normalize_line(_p_element_text(para._p))
                    if text:
                        segments.append(text)
                        seg = {
                            "kind": kind,
                            "path": path + [[nt_idx, r_idx, c_idx]],
                            "para": p_idx,
                        }
                        if part is not None:
                            seg["part"] = part
                        structure["segments"].append(seg)
                        if paras_out is not None:
                            paras_out.append(text)
                _emit_nested_cell_paras(
                    ncell, path + [[nt_idx, r_idx, c_idx]],
                    kind, part, segments, structure, paras_out)


def _resolve_nested_cell(tables, path):
    """Ячейка по цепочке координат вложенности (см. _emit_nested_cell_paras)."""
    t, r, c = path[0]
    cell = tables[t].rows[r].cells[c]
    for t2, r2, c2 in path[1:]:
        cell = _cell_tables(cell)[t2].rows[r2].cells[c2]
    return cell


# Namespace расширенных свойств комментариев Word (people.xml и пр.)
_W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"


def _comments_element(doc):
    """Корень word/comments.xml или None, если части комментариев нет."""
    for part in doc.part.package.iter_parts():
        if str(part.partname) == "/word/comments.xml":
            try:
                return part.element
            except AttributeError:
                return None
    return None


def read_anon_marker(path) -> Optional[str]:
    """Маркер версии анонимизатора из свойств копии (DOCX/XLSX).

    None — маркера нет (копия создана старой версией или файл другого
    формата): такую копию переиспользовать нельзя.
    """
    ext = Path(path).suffix.lower()
    try:
        if ext == ".docx":
            props = Document(str(path)).core_properties
            return props.comments
        if ext == ".xlsx":
            wb = load_workbook(str(path), read_only=True)
            try:
                return wb.properties.keywords
            finally:
                wb.close()
    except Exception:  # noqa: BLE001 - нечитаемая копия = переиспользовать нельзя
        return None
    return None


def _scrub_people_xml(docx_bytes: bytes, author_tokens: dict) -> tuple[bytes, int]:
    """Заменить имена авторов комментариев в word/people.xml их токенами.

    Имена авторов живут не только в w:author/comments.xml, но и в
    word/people.xml (w15:person) — без этого авторы утечут из
    анонимизированной копии. Возвращает (новые байты пакета, число замен).
    """
    if not author_tokens:
        return docx_bytes, 0
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zin:
        infos = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in infos}
    name = "word/people.xml"
    if name not in data:
        return docx_bytes, 0
    try:
        root = _lxml.fromstring(data[name])
    except Exception:  # noqa: BLE001 - битый people.xml не роняет сборку
        return docx_bytes, 0
    changed = 0
    for person in root.iter(f"{{{_W15_NS}}}person"):
        val = person.get(f"{{{_W15_NS}}}author")
        if val in author_tokens:
            person.set(f"{{{_W15_NS}}}author", author_tokens[val])
            changed += 1
        # userId в presenceInfo часто равен имени автора (логину) — тоже PII
        for presence in person.iter(f"{{{_W15_NS}}}presenceInfo"):
            uid = presence.get(f"{{{_W15_NS}}}userId")
            if uid in author_tokens:
                presence.set(f"{{{_W15_NS}}}userId", author_tokens[uid])
                changed += 1
    if not changed:
        return docx_bytes, 0
    data[name] = _lxml.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, data[i.filename])
    return buf.getvalue(), changed


def _set_paragraph_text(para, text: str) -> None:
    """Заменить текст абзаца (включая текст гиперссылок), сохранив сам
    абзац и его стиль. Неизменённые абзацы не затрагиваются. Текст внутри
    вложенных фигур не трогается — у него свои сегменты."""
    p_el = para._p
    if _normalize_line(_p_element_text(p_el)) == text:
        return
    own_box = _nearest_txbx(p_el)

    # Поля Word (PAGE, NUMPAGES, перекрёстные ссылки...): в изменённом
    # абзаце структура полей удаляется. Word вычисляет значения полей при
    # рендере и дорисовывает их поверх анонимизированного текста — футер
    # «Страница 2 из 27» после маскировки превращается в «[LOC_11]227»
    # (токен + «живые» номера страницы). Неизменённые абзацы не затрагиваются
    # (ранний выход выше) — поля там работают как раньше, а де-анонимизация
    # возвращает исходный текст, и поля оживают снова.
    for r in list(p_el.iter(qn("w:r"))):
        if _nearest_txbx(r) is not own_box:
            continue
        if (r.find(qn("w:fldChar")) is not None
                or r.find(qn("w:instrText")) is not None):
            parent = r.getparent()
            if parent is not None:
                parent.remove(r)
    for fs in list(p_el.iter(qn("w:fldSimple"))):
        if _nearest_txbx(fs) is own_box:
            parent = fs.getparent()
            if parent is not None:
                parent.remove(fs)

    targets = [
        t for t in p_el.iter(qn("w:t"))
        if _nearest_txbx(t) is own_box
    ]
    if targets:
        for t in targets:
            t.text = ""
        first = targets[0]
        first.text = text
        first.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    else:
        para.add_run(text)


# ==================== DOCX: колонтитулы, фигуры, картинки ====================

# Теги с изображениями (Clark-нотация; префиксы nsmap python-docx не
# гарантирует для pic/a/v, поэтому используем явные URI)
_PIC_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"
_BLIP_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_IMAGEDATA_TAG = "{urn:schemas-microsoft-com:vml}imagedata"

# Удалять картинки (логотипы) из колонтитулов анонимизированных DOCX.
# Отключается REMOVE_HEADER_IMAGES=0
REMOVE_HF_IMAGES = os.getenv(
    "REMOVE_HEADER_IMAGES", "1"
).strip().lower() in ("1", "true", "yes", "on")


def _docx_hf_objects(doc):
    """Уникальные части колонтитулов DOCX в детерминированном порядке.

    Пропускаются колонтитулы «как в предыдущем разделе» (нет собственной
    части) и повторные ссылки разных секций на одну и ту же часть — иначе
    сегменты задвоятся и нарушится соответствие «строка текста <-> сегмент»,
    на котором строится обратная сборка.
    """
    seen = set()
    result = []
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            try:
                if hf.is_linked_to_previous:
                    continue
            except Exception:  # noqa: BLE001 - битая ссылка не должна ронять парсинг
                continue
            partname = str(hf.part.partname)
            if partname in seen:
                continue
            seen.add(partname)
            result.append(hf)
    return result


def _top_txbx_contents(root):
    """Корневые w:txbxContent (без вложенных) в порядке следования.

    Вложенные фигуры обрабатываются вместе с внешним контейнером, чтобы
    каждый абзац попал в сегменты ровно один раз.
    """
    result = []
    for el in root.iter(qn("w:txbxContent")):
        if any(True for _ in el.iterancestors(qn("w:txbxContent"))):
            continue
        result.append(el)
    return result


def _txbx_para_elements(root):
    """Для каждого корневого w:txbxContent — список его элементов w:p.

    Используется и при парсинге, и при сборке — порядок индексов обязан
    совпадать (детерминированный обход lxml).
    """
    return [list(content.iter(qn("w:p"))) for content in _top_txbx_contents(root)]


def _strip_hf_images(doc) -> int:
    """Удалить картинки из всех частей колонтитулов.

    Затрагиваются только узлы с изображениями: w:drawing, содержащие
    pic:pic или a:blip, и w:pict с v:imagedata. Фигуры-надписи без
    картинок не удаляются — их текст уже анонимизирован.
    Попутно удаляются Relationship-записи удалённых картинок, иначе
    media-файлы остаются в пакете (связь делает их «используемыми»).
    """
    removed = 0
    for hf in _docx_hf_objects(doc):
        root = hf.part.element
        targets = []  # собираем заранее: удаление во время iter небезопасно
        rids = set()  # rId картинок для последующей чистки .rels части
        for el in root.iter():
            if el.tag == qn("w:drawing"):
                if (el.find(f".//{_PIC_TAG}") is not None
                        or el.find(f".//{_BLIP_TAG}") is not None):
                    targets.append(el)
            elif el.tag == qn("w:pict"):
                if el.find(f".//{_IMAGEDATA_TAG}") is not None:
                    targets.append(el)
        for el in targets:
            for node in el.iter():
                for attr in (qn("r:embed"), qn("r:id"), qn("r:link")):
                    rid = node.get(attr)
                    if rid:
                        rids.add(rid)
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
                removed += 1
        for rid in rids:
            try:
                hf.part.drop_rel(rid)
            except (KeyError, AttributeError):
                # Связи может не быть (external) либо старая версия python-docx
                logger.debug("Не удалось удалить связь %s части %s",
                             rid, hf.part.partname)
    return removed


def _drop_orphan_media(docx_bytes: bytes) -> tuple[bytes, int]:
    """Вымыть из DOCX-пакета media-файлы без ссылок в .rels.

    После удаления картинок из колонтитулов их байты остаются лежать в
    word/media/ — это утечка (логотипы читаются из ZIP напрямую).
    Попутно чистятся осиротевшие Relationship-записи, чтобы пакет
    оставался валидным. Возвращает (новые байты, число удалённых файлов).
    """
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zin:
        infos = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in infos}

    def resolve(base_dir: str, target: str) -> str:
        return posixpath.normpath(posixpath.join(base_dir, target)).lstrip("/")

    # 1) Какие части реально используются (все .rels, только Internal)
    used = set()
    for name, blob in data.items():
        if not name.endswith(".rels"):
            continue
        base_dir = posixpath.dirname(posixpath.dirname(name))
        root = ET.fromstring(blob)
        for rel in root:
            if (rel.get("TargetMode") or "Internal").lower() == "external":
                continue
            used.add(resolve(base_dir, rel.get("Target", "")))

    dropped = {
        n for n in data
        if n.startswith("word/media/") and n not in used
    }
    if not dropped:
        return docx_bytes, 0

    # 2) Чистим .rels от связей с удалёнными частями
    for name in list(data):
        if not name.endswith(".rels"):
            continue
        base_dir = posixpath.dirname(posixpath.dirname(name))
        root = ET.fromstring(data[name])
        changed = False
        for rel in list(root):
            tmode = (rel.get("TargetMode") or "Internal").lower()
            resolved = (rel.get("Target", "") if tmode == "external"
                        else resolve(base_dir, rel.get("Target", "")))
            if resolved in dropped:
                root.remove(rel)
                changed = True
        if changed:
            data[name] = ET.tostring(
                root, xml_declaration=True, encoding="UTF-8")

    # 3) Перезаписываем пакет без удалённых media
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            if i.filename in dropped:
                continue
            zout.writestr(i, data[i.filename])
    return buf.getvalue(), len(dropped)


def _scrub_hyperlink_targets(docx_bytes: bytes) -> tuple[bytes, int]:
    """Заменить внешние цели гиперссылок в *.rels пакета на нейтральное «#».

    Адрес сайта/почты может храниться не только в тексте (его маскирует
    NER-слой), но и в Relationship-целях гиперссылок: тогда он утечёт из
    анонимизированной копии даже при замаскированном тексте ссылки.
    Возвращает (новые байты пакета, число обезличенных ссылок).
    Вызывается только при АНОНИМИЗИРУЮЩЕЙ сборке (не при де-анонимизации).
    """
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zin:
        infos = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in infos}

    changed = 0
    for name, blob in data.items():
        if not name.endswith(".rels"):
            continue
        try:
            root = ET.fromstring(blob)
        except ET.ParseError:
            continue
        modified = False
        for rel in root:
            if (rel.get("TargetMode") or "Internal").lower() != "external":
                continue
            if "hyperlink" not in (rel.get("Type") or "").lower():
                continue
            target = rel.get("Target") or ""
            if re.match(r"(?i)^(https?://|ftp://|mailto:|www\.)", target):
                rel.set("Target", "#")
                modified = True
                changed += 1
        if modified:
            data[name] = ET.tostring(
                root, xml_declaration=True, encoding="UTF-8")

    if not changed:
        return docx_bytes, 0

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, data[i.filename])
    return buf.getvalue(), changed


# Версия экстрактора/сборщика. Записывается маркером в свойства
# анонимизированной копии (DOCX: docProps/core.xml → comments;
# XLSX: keywords). Повторная анонимизация по неизменённому исходнику
# переиспользует копию ТОЛЬКО при совпадении версии — копии, созданные
# другой версией парсера (например, до добавления поддержки вложенных
# таблиц/комментариев), автоматически пересоздаются.
PARSER_VERSION = "3"
ANON_MARKER_PREFIX = "anonymizer_proxy/parser:"
ANON_MARKER = ANON_MARKER_PREFIX + PARSER_VERSION


class FileParser:
    """Парсер файлов различных форматов"""
    
    SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".xml", ".txt", ".md"}
    
    @classmethod
    def is_supported(cls, filename: str) -> bool:
        """Проверить, поддерживается ли формат файла"""
        ext = Path(filename).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS
    
    @classmethod
    def get_content_type(cls, filename: str) -> str:
        """Определить тип контента по расширению"""
        ext = Path(filename).suffix.lower()
        content_types = {
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ".xml": "application/xml",
            ".txt": "text/plain",
            ".md": "text/markdown",
        }
        return content_types.get(ext, "application/octet-stream")
    
    async def parse(self, content: bytes, filename: str) -> ParsedContent:
        """
        Распарсить содержимое файла
        
        Args:
            content: Бинарное содержимое файла
            filename: Имя файла (для определения формата)
        
        Returns:
            ParsedContent с извлечённым текстом и структурой
        """
        ext = Path(filename).suffix.lower()
        
        if ext == ".docx":
            return self._parse_docx(content)
        elif ext == ".xlsx":
            return self._parse_xlsx(content)
        elif ext == ".xml":
            return self._parse_xml(content)
        elif ext in (".txt", ".md"):
            return self._parse_text(content)
        else:
            raise ValueError(f"Неподдерживаемый формат файла: {ext}")
    
    def _parse_docx(self, content: bytes) -> ParsedContent:
        """Парсинг DOCX файлов с сохранением структуры абзацев и таблиц.

        Каждый абзац тела документа и каждый абзац ячейки таблицы дают
        ровно одну строку текста (сегмент); координаты сегментов
        сохраняются в structure["segments"] для обратной сборки.
        """
        doc = Document(io.BytesIO(content))

        segments = []
        structure = {"paragraphs": [], "tables": [], "segments": []}

        # Извлекаем параграфы тела документа
        for i, para in enumerate(doc.paragraphs):
            text = _normalize_line(_p_element_text(para._p))
            if text:
                segments.append(text)
                structure["segments"].append({"kind": "para", "para": i})
                structure["paragraphs"].append({
                    "index": i,
                    "text": text,
                    "style": para.style.name if para.style else None
                })

        # Извлекаем таблицы: каждый абзац ячейки — отдельный сегмент
        # (текст ячейки может состоять из нескольких абзацев)
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row_idx, row in enumerate(table.rows):
                row_data = []
                for cell_idx, cell in enumerate(row.cells):
                    cell_paras = []
                    for para_idx, para in enumerate(cell.paragraphs):
                        text = _normalize_line(_p_element_text(para._p))
                        if text:
                            segments.append(text)
                            structure["segments"].append({
                                "kind": "cell_para",
                                "table": table_idx,
                                "row": row_idx,
                                "cell": cell_idx,
                                "para": para_idx
                            })
                            cell_paras.append(text)
                    row_data.append("\n".join(cell_paras))
                table_data.append(row_data)
            structure["tables"].append({
                "table_index": table_idx,
                "data": table_data
            })

        # Вложенные таблицы (w:tbl внутри ячеек): doc.tables видит только
        # таблицы верхнего уровня тела, cell.paragraphs не спускается во
        # вложенные — титульные листы («УТВЕРЖДАЮ» с ООО «Газпром нефтехим
        # Салават» и подписантами) оставались неанонимизированными
        # (багрепорт 2026-09-08).
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    _emit_nested_cell_paras(
                        cell, [[t_idx, r_idx, c_idx]],
                        "nested_cell_para", None, segments, structure)

        # Колонтитулы: абзацы и ячейки таблиц каждой уникальной части
        structure["headers_footers"] = []
        for hf in _docx_hf_objects(doc):
            pn = str(hf.part.partname)
            title = Path(pn).name  # например header1.xml
            entry = {
                "part": pn,
                "title": title,
                "label": "Верхний колонтитул" if "header" in title.lower()
                         else "Нижний колонтитул",
                "paras": [],
            }
            for i, para in enumerate(hf.paragraphs):
                text = _normalize_line(_p_element_text(para._p))
                if text:
                    segments.append(text)
                    structure["segments"].append(
                        {"kind": "hf_para", "part": pn, "para": i})
                    entry["paras"].append(text)
            for ti, table in enumerate(hf.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, para in enumerate(cell.paragraphs):
                            text = _normalize_line(_p_element_text(para._p))
                            if text:
                                segments.append(text)
                                structure["segments"].append({
                                    "kind": "hf_cell_para",
                                    "part": pn,
                                    "table": ti, "row": ri,
                                    "cell": ci, "para": pi,
                                })
                                entry["paras"].append(text)
                        _emit_nested_cell_paras(
                            cell, [[ti, ri, ci]],
                            "nested_hf_cell_para", pn,
                            segments, structure, entry["paras"])
            if entry["paras"]:
                structure["headers_footers"].append(entry)

        # Текстовые фигуры (w:txbxContent): тело документа и колонтитулы.
        # Дубли в mc:Choice/mc:Fallback дают отдельные сегменты — замена
        # согласована, т.к. маппинг value->token глобальный по сессии.
        structure["textboxes"] = []
        hosts = [("document", doc.element)]
        for hf in _docx_hf_objects(doc):
            hosts.append((str(hf.part.partname), hf.part.element))
        for host_name, host_root in hosts:
            for k, p_elements in enumerate(_txbx_para_elements(host_root)):
                shape_entry = {"host": host_name, "index": k, "paras": []}
                for j, p_el in enumerate(p_elements):
                    text = _normalize_line(_p_element_text(p_el))
                    if text:
                        segments.append(text)
                        structure["segments"].append({
                            "kind": "txbx_para",
                            "host": host_name,
                            "txbx": k, "para": j,
                        })
                        shape_entry["paras"].append(text)
                if shape_entry["paras"]:
                    structure["textboxes"].append(shape_entry)

        # Комментарии (word/comments.xml): ни текст, ни авторы комментариев
        # не попадают ни в тело, ни в python-docx-абзацы — раньше они
        # неанонимизированными уходили в облако (багрепорт 2026-09-08:
        # автор Sasha, текст комментария с ФИО). Автор — атрибут w:author:
        # добавляем его отдельным сегментом, чтобы NER увидел значение, а
        # сборка записала токен обратно в атрибут.
        structure["comments"] = []
        comments_root = _comments_element(doc)
        if comments_root is not None:
            for c_idx, c_el in enumerate(comments_root.findall(qn("w:comment"))):
                entry = {"index": c_idx, "author": None, "paras": []}
                author = (c_el.get(qn("w:author")) or "").strip()
                if author:
                    segments.append(author)
                    structure["segments"].append(
                        {"kind": "comment_author", "comment": c_idx})
                    entry["author"] = author
                for p_idx, p_el in enumerate(c_el.findall(qn("w:p"))):
                    text = _normalize_line(_p_element_text(p_el))
                    if text:
                        segments.append(text)
                        structure["segments"].append({
                            "kind": "comment_para",
                            "comment": c_idx, "para": p_idx,
                        })
                        entry["paras"].append(text)
                if entry["author"] or entry["paras"]:
                    structure["comments"].append(entry)

        # Markdown-представление с таблицами (для review и облака):
        # абзацы и таблицы в порядке документа, таблицы -> markdown-таблицы
        # (включая вложенные); затем блоки колонтитулов, фигур и комментариев
        markdown_parts = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                text = _normalize_line(_p_element_text(child))
                if text:
                    markdown_parts.append(text)
            elif child.tag == qn("w:tbl"):
                table = Table(child, doc)
                markdown_parts.append(_table_to_markdown_deep(table))

        extra_md = []
        for entry in structure["headers_footers"]:
            extra_md.append(
                f"**[{entry['label']} {entry['title']}]**\n"
                + "\n".join(entry["paras"]))
        for entry in structure["textboxes"]:
            extra_md.append(
                f"**[Фигура {entry['index']} ({entry['host']})]**\n"
                + "\n".join(entry["paras"]))
        for entry in structure["comments"]:
            lines = [f"**[Комментарий {entry['index'] + 1}]**"]
            if entry["author"]:
                lines.append(f"Автор: {entry['author']}")
            lines.extend(entry["paras"])
            extra_md.append("\n".join(lines))
        if extra_md:
            markdown_parts.extend(extra_md)

        markdown = "\n\n".join(p for p in markdown_parts if p)

        return ParsedContent(
            text="\n".join(segments),
            structure=structure,
            metadata={
                "format": "docx",
                "paragraphs_count": len(doc.paragraphs),
                "tables_count": len(doc.tables),
                "headers_footers_count": len(structure["headers_footers"]),
                "textboxes_count": len(structure["textboxes"]),
                "comments_count": len(structure["comments"]),
            },
            markdown=markdown,
        )

    def _parse_xlsx(self, content: bytes) -> ParsedContent:
        """Парсинг XLSX файлов с сохранением структуры листов и ячеек.

        Ячейки читаются как есть (data_only=False): формулы остаются
        формулами — при обратной сборке неизменённые ячейки (формулы,
        числа, даты) не перезаписываются. Координаты непустых ячеек
        сохраняются в structure["segments"].
        """
        wb = load_workbook(io.BytesIO(content))

        segments = []
        structure = {"sheets": [], "segments": []}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = {"name": sheet_name, "rows": [], "comments": []}

            for row in ws.iter_rows():
                row_values = []
                for cell in row:
                    cell_str = _xlsx_cell_text(cell)
                    if cell_str:
                        row_values.append(cell_str)
                        segments.append(cell_str)
                        structure["segments"].append({
                            "kind": "cell",
                            "sheet": sheet_name,
                            "row": cell.row,
                            "col": cell.column
                        })
                    else:
                        row_values.append("")

                    # Комментарий ячейки (примечание Excel): текст и автор
                    # не видны в значениях ячеек — раньше не анонимизировались
                    # (примечания поддерживаются openpyxl; тредовые комментарии
                    # нового Excel openpyxl не читает — ограничение)
                    comment = cell.comment
                    if comment is not None:
                        c_text = _normalize_line(comment.text or "")
                        if c_text:
                            segments.append(c_text)
                            structure["segments"].append({
                                "kind": "xlsx_comment",
                                "sheet": sheet_name,
                                "row": cell.row,
                                "col": cell.column,
                            })
                            sheet_data["comments"].append({
                                "ref": cell.coordinate,
                                "author": (comment.author or "").strip(),
                                "text": c_text,
                            })
                        author = (comment.author or "").strip()
                        if author:
                            segments.append(author)
                            structure["segments"].append({
                                "kind": "xlsx_comment_author",
                                "sheet": sheet_name,
                                "row": cell.row,
                                "col": cell.column,
                            })
                sheet_data["rows"].append(row_values)

            structure["sheets"].append(sheet_data)

        # Markdown-представление: каждый лист -> markdown-таблица
        # (дисплейное представление через _fmt_cell_value: даты ISO,
        # без _x000D_ — сегменты text остаются через _cell_to_line);
        # комментарии ячеек — отдельными блоками после таблицы листа
        markdown_parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows():
                cells = [_fmt_cell_value(cell.value) for cell in row]
                if not any(cells):
                    continue
                rows.append(cells)
            markdown_parts.append(f"## {sheet_name}\n\n{_rows_to_markdown(rows)}")
            sheet_entry = next(
                (s for s in structure["sheets"] if s["name"] == sheet_name),
                None)
            for c in (sheet_entry or {}).get("comments", []):
                lines = [f"**[Комментарий {sheet_name}!{c['ref']}]**"]
                if c.get("author"):
                    lines.append(f"Автор: {c['author']}")
                lines.append(c["text"])
                markdown_parts.append("\n".join(lines))
        markdown = "\n\n".join(markdown_parts)

        return ParsedContent(
            text="\n".join(segments),
            structure=structure,
            metadata={"format": "xlsx", "sheets_count": len(wb.sheetnames)},
            markdown=markdown,
        )

    def _parse_xml(self, content: bytes) -> ParsedContent:
        """
        Парсинг XML файлов (включая MS Project XML)
        Извлекает текстовое содержимое, сохраняя структуру тегов
        """
        try:
            root = ET.fromstring(content)
        except ET.ParseError as e:
            raise ValueError(f"Ошибка парсинга XML: {e}")
        
        text_parts = []
        structure = {"elements": []}
        
        # Рекурсивно извлекаем текст из всех элементов
        def extract_text(element, path=""):
            current_path = f"{path}/{element.tag}" if path else element.tag
            
            # Текст элемента
            if element.text and element.text.strip():
                text = element.text.strip()
                text_parts.append(text)
                structure["elements"].append({
                    "path": current_path,
                    "text": text,
                    "type": "text"
                })
            
            # Атрибуты (могут содержать имена, названия)
            for attr_name, attr_value in element.attrib.items():
                if attr_value and len(attr_value) > 2:  # Игнорируем короткие значения
                    # Проверяем, может ли атрибут содержать PII
                    pii_attrs = ["name", "Name", "ResourceName", "Author", "Manager"]
                    if any(pii in attr_name for pii in pii_attrs):
                        text_parts.append(attr_value)
                        structure["elements"].append({
                            "path": f"{current_path}@{attr_name}",
                            "text": attr_value,
                            "type": "attribute"
                        })
            
            # Рекурсия для дочерних элементов
            for child in element:
                extract_text(child, current_path)
            
            # Tail текст (текст после закрывающего тега)
            if element.tail and element.tail.strip():
                tail_text = element.tail.strip()
                text_parts.append(tail_text)
                structure["elements"].append({
                    "path": f"{current_path}/tail",
                    "text": tail_text,
                    "type": "tail"
                })
        
        extract_text(root)
        
        full_text = "\n".join(text_parts)
        
        # Определяем тип XML (MS Project или другой)
        xml_type = "generic"
        if root.tag.endswith("Project") or "Project" in root.tag:
            xml_type = "msproject"
        
        return ParsedContent(
            text=full_text,
            structure=structure,
            metadata={
                "format": "xml",
                "xml_type": xml_type,
                "root_tag": root.tag,
                "elements_count": len(structure["elements"])
            },
            markdown=full_text,
        )
    
    def _parse_text(self, content: bytes) -> ParsedContent:
        """Парсинг текстовых файлов"""
        # Пытаемся определить кодировку
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")  # Windows-1251 для русского
            except UnicodeDecodeError:
                text = content.decode("utf-8", errors="replace")
        
        return ParsedContent(
            text=text,
            structure={"type": "plain_text"},
            metadata={"format": "text", "encoding": "auto-detected"},
            markdown=text,
        )


class FileAssembler:
    """Сборщик файлов обратно из анонимизированного текста"""
    
    async def assemble(
        self,
        original_content: bytes,
        original_filename: str,
        anonymized_text: str,
        structure: dict,
        strip_hf_images: bool = True,
        scrub_metadata: bool = True,
    ) -> bytes:
        """
        Собрать файл обратно с анонимизированным содержимым

        Args:
            original_content: Оригинальное содержимое файла
            original_filename: Имя файла
            anonymized_text: Анонимизированный текст
            structure: Структура из ParsedContent
            strip_hf_images: Удалять картинки колонтитулов DOCX (логотипы);
                при де-анонимизации передавайте False — восстановленный
                файл должен остаться с оригинальным оформлением
            scrub_metadata: Обезличивать метаданные пакета DOCX
                (автор, компания, HLinks); False — для де-анонимизации

        Returns:
            Бинарное содержимое файла с заменённым текстом
        """
        ext = Path(original_filename).suffix.lower()
        
        if ext == ".docx":
            return self._assemble_docx(
                original_content, anonymized_text, structure,
                strip_hf_images=strip_hf_images,
                scrub_metadata=scrub_metadata,
            )
        elif ext == ".xlsx":
            return self._assemble_xlsx(
                original_content, anonymized_text, structure,
                scrub_metadata=scrub_metadata,
            )
        elif ext == ".xml":
            return self._assemble_xml(original_content, anonymized_text, structure)
        elif ext in (".txt", ".md"):
            return anonymized_text.encode("utf-8")
        else:
            raise ValueError(f"Неподдерживаемый формат файла: {ext}")
    
    def _assemble_docx(
        self,
        original_content: bytes,
        anonymized_text: str,
        structure: dict,
        strip_hf_images: bool = True,
        scrub_metadata: bool = True,
    ) -> bytes:
        """Сборка DOCX файла: текст записывается обратно по координатам,
        сохранённым при парсинге (абзацы тела, ячейки таблиц, колонтитулы,
        текстовые фигуры). Сетка таблиц и стили абзацев сохраняются;
        неизменённые абзацы не затрагиваются.
        Из колонтитулов удаляются картинки (логотипы), а их media-файлы
        вымываются из пакета (см. REMOVE_HEADER_IMAGES)."""
        doc = Document(io.BytesIO(original_content))
        parts = anonymized_text.split("\n")
        segments = structure.get("segments", [])

        # Быстрый доступ к частям колонтитулов по имени части
        hf_by_part = {}
        hosts = {"document": doc.element}
        for hf in _docx_hf_objects(doc):
            pn = str(hf.part.partname)
            hf_by_part[pn] = hf
            hosts[pn] = hf.part.element

        # Часть комментариев (python-docx её сериализует при doc.save)
        comments_root = _comments_element(doc)
        comment_elems = (
            comments_root.findall(qn("w:comment"))
            if comments_root is not None else []
        )
        author_tokens: dict = {}  # исходный автор -> токен (для people.xml)

        for seg, part in zip(segments, parts):
            kind = seg["kind"]
            if kind == "para":
                para = doc.paragraphs[seg["para"]]
            elif kind == "cell_para":
                cell = doc.tables[seg["table"]].rows[seg["row"]].cells[seg["cell"]]
                para = cell.paragraphs[seg["para"]]
            elif kind == "nested_cell_para":
                try:
                    cell = _resolve_nested_cell(doc.tables, seg["path"])
                    para = cell.paragraphs[seg["para"]]
                except IndexError:
                    logger.warning(
                        "Пропущен сегмент вложенной таблицы: путь %s вне "
                        "диапазона", seg.get("path"))
                    continue
            elif kind == "hf_para":
                hf = hf_by_part.get(seg["part"])
                if hf is None:
                    logger.warning(
                        "Пропущен сегмент колонтитула %s: часть недоступна",
                        seg["part"])
                    continue
                para = hf.paragraphs[seg["para"]]
            elif kind == "hf_cell_para":
                hf = hf_by_part.get(seg["part"])
                if hf is None:
                    logger.warning(
                        "Пропущен сегмент колонтитула %s: часть недоступна",
                        seg["part"])
                    continue
                cell = hf.tables[seg["table"]].rows[seg["row"]].cells[seg["cell"]]
                para = cell.paragraphs[seg["para"]]
            elif kind == "nested_hf_cell_para":
                hf = hf_by_part.get(seg.get("part"))
                if hf is None:
                    logger.warning(
                        "Пропущен сегмент вложенной таблицы колонтитула %s: "
                        "часть недоступна", seg.get("part"))
                    continue
                try:
                    cell = _resolve_nested_cell(hf.tables, seg["path"])
                    para = cell.paragraphs[seg["para"]]
                except IndexError:
                    logger.warning(
                        "Пропущен сегмент вложенной таблицы колонтитула: "
                        "путь %s вне диапазона", seg.get("path"))
                    continue
            elif kind == "txbx_para":
                boxes = _txbx_para_elements(hosts.get(seg["host"], doc.element))
                try:
                    p_el = boxes[seg["txbx"]][seg["para"]]
                except IndexError:
                    logger.warning(
                        "Пропущен сегмент фигуры (%s/%s): индекс вне диапазона",
                        seg["host"], seg["txbx"])
                    continue
                para = Paragraph(p_el, None)
            elif kind == "comment_author":
                try:
                    comment_elems[seg["comment"]].set(qn("w:author"), part)
                except IndexError:
                    logger.warning(
                        "Пропущен автор комментария %s: индекс вне диапазона",
                        seg["comment"])
                    continue
                continue  # атрибут, не абзац — _set_paragraph_text не нужен
            elif kind == "comment_para":
                try:
                    p_el = (comment_elems[seg["comment"]]
                            .findall(qn("w:p"))[seg["para"]])
                except IndexError:
                    logger.warning(
                        "Пропущен текст комментария %s/%s: индекс вне "
                        "диапазона", seg["comment"], seg["para"])
                    continue
                para = Paragraph(p_el, None)
            else:
                logger.warning("Неизвестный тип сегмента: %s", kind)
                continue
            _set_paragraph_text(para, part)

        # Запоминаем соответствие «исходный автор комментария -> токен»
        # для чистки word/people.xml (имена авторов дублируются там)
        for seg, part in zip(segments, parts):
            if seg.get("kind") == "comment_author":
                for entry in structure.get("comments", []):
                    if (entry.get("index") == seg.get("comment")
                            and entry.get("author")):
                        author_tokens[entry["author"]] = part

        if scrub_metadata:
            _scrub_docx_metadata(doc)
            # Маркер версии парсера: повторная анонимизация по неизменённому
            # исходнику переиспользует копию только при совпадении версии
            doc.core_properties.comments = ANON_MARKER
        else:
            # финальный (де-анонимизированный) файл не несёт служебный маркер
            doc.core_properties.comments = ""

        # Удаляем картинки (логотипы) из колонтитулов и вымываем их
        # байты из пакета — иначе логотипы остаются в word/media/
        do_strip = strip_hf_images and REMOVE_HF_IMAGES
        removed_images = 0
        if do_strip:
            removed_images = _strip_hf_images(doc)

        # Сохраняем в bytes
        output = io.BytesIO()
        doc.save(output)
        result = output.getvalue()

        if scrub_metadata and author_tokens:
            # Имена авторов комментариев дублируются в word/people.xml —
            # без чистки они утечут из анонимизированной копии
            result, scrubbed_people = _scrub_people_xml(result, author_tokens)
            if scrubbed_people:
                logger.info(
                    "Обезличено авторов комментариев в people.xml: %d",
                    scrubbed_people)

        if scrub_metadata:
            # Внешние цели гиперссылок (адреса сайтов, mailto:) не должны
            # оставаться в анонимизированной копии даже при замаскированном
            # тексте ссылки
            result, scrubbed_links = _scrub_hyperlink_targets(result)
            if scrubbed_links:
                logger.info(
                    "Обезличено внешних гиперссылок в пакете: %d",
                    scrubbed_links)

        if do_strip and removed_images:
            result, dropped = _drop_orphan_media(result)
            logger.info(
                "Удалено картинок из колонтитулов: %s (media-файлов из "
                "пакета: %s)", removed_images, dropped)
        return result

    def _assemble_xlsx(self, original_content: bytes, anonymized_text: str, structure: dict, scrub_metadata: bool = True) -> bytes:
        """Сборка XLSX файла: значения записываются обратно по координатам,
        сохранённым при парсинге. Ячейки, чей текст не изменился, не
        затрагиваются — это сохраняет формулы, числа и даты. Комментарии
        ячеек (текст и автор) восстанавливаются по своим координатам."""
        wb = load_workbook(io.BytesIO(original_content))
        parts = anonymized_text.split("\n")
        segments = structure.get("segments", [])

        for seg, part in zip(segments, parts):
            if seg["kind"] in ("xlsx_comment", "xlsx_comment_author"):
                ws = wb[seg["sheet"]]
                cell = ws.cell(row=seg["row"], column=seg["col"])
                comment = cell.comment
                if comment is None:
                    logger.warning(
                        "Пропущен комментарий %s!%s: ячейка без комментария",
                        seg["sheet"], cell.coordinate)
                    continue
                if seg["kind"] == "xlsx_comment_author":
                    if (comment.author or "") != part:
                        comment.author = part
                elif _normalize_line(comment.text or "") != part:
                    comment.text = part
                continue
            ws = wb[seg["sheet"]]
            cell = ws.cell(row=seg["row"], column=seg["col"])
            if _xlsx_cell_text(cell) != part:
                cell.value = _restore_cell_value(cell.value, part)

        output = io.BytesIO()
        wb.save(output)
        result = output.getvalue()

        if scrub_metadata:
            # Маркер версии парсера в свойствах копии (см. PARSER_VERSION)
            try:
                wb.properties.keywords = ANON_MARKER
                buf = io.BytesIO()
                wb.save(buf)
                result = buf.getvalue()
            except Exception as exc:  # noqa: BLE001
                logger.warning("Не удалось записать маркер версии: %s", exc)
        return result

    def _assemble_xml(self, original_content: bytes, anonymized_text: str, structure: dict) -> bytes:
        """Сборка XML файла"""
        root = ET.fromstring(original_content)
        
        # Разбиваем анонимизированный текст на части
        anon_parts = anonymized_text.split("\n")
        part_idx = 0
        
        # Создаём маппинг path -> новый текст
        replacements = {}
        for elem_info in structure.get("elements", []):
            if part_idx < len(anon_parts):
                replacements[elem_info["path"]] = anon_parts[part_idx]
                part_idx += 1
        
        # Рекурсивно заменяем текст
        def replace_text(element, path=""):
            current_path = f"{path}/{element.tag}" if path else element.tag
            
            if element.text and element.text.strip():
                if current_path in replacements:
                    element.text = replacements[current_path]
            
            for attr_name, attr_value in element.attrib.items():
                attr_path = f"{current_path}@{attr_name}"
                if attr_path in replacements:
                    element.set(attr_name, replacements[attr_path])
            
            for child in element:
                replace_text(child, current_path)
            
            if element.tail and element.tail.strip():
                tail_path = f"{current_path}/tail"
                if tail_path in replacements:
                    element.tail = replacements[tail_path]
        
        replace_text(root)
        
        # ET не поддерживает pretty_print, используем indent для форматирования
        ET.indent(root, space="  ")
        return ET.tostring(root, encoding="unicode").encode("utf-8")
