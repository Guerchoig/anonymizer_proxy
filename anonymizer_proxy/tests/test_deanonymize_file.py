"""
Функциональные тесты де-анонимизации файла (/api/deanonymize_file).

Проверяют замену плейсхолдеров на реальные значения:
1. .md/.txt — текстовая замена.
2. .docx — с сохранением структуры (через FileParser + FileAssembler).

Запуск: python anonymizer_proxy\\tests\\test_deanonymize_file.py (из корня проекта)
"""
import asyncio
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from anonymizer_proxy.proxy.handlers import RequestHandler
from anonymizer_proxy.tests.test_tools_manual import FakeNER, FakeStore, FakeOpenRouter


def make_handler(store):
    return RequestHandler(
        ner_service=FakeNER({}),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )


async def test_deanonymize_file_md():
    store = FakeStore()
    await store.add_mapping("s1", "Ивана Петрова", "PERSON")
    handler = make_handler(store)

    with tempfile.NamedTemporaryFile(
        suffix=".md", mode="w", encoding="utf-8", delete=False
    ) as f:
        f.write("# Заголовок\n\nПроверка документа от [PERSON_1]\n")
        path = Path(f.name)

    try:
        result = await handler.handle_deanonymize_file("s1", str(path))
        text = path.read_text(encoding="utf-8")
        assert "[PERSON_1]" not in text, text
        assert "Ивана Петрова" in text, text
        assert result["file_path"] == str(path)
        print("TEST 1 OK: deanonymize_file .md — плейсхолдер заменён значением")
    finally:
        path.unlink(missing_ok=True)


async def test_deanonymize_file_docx():
    from docx import Document

    store = FakeStore()
    await store.add_mapping("s1", "Ивана Петрова", "PERSON")
    handler = make_handler(store)

    doc = Document()
    doc.add_paragraph("Документ подготовлен для [PERSON_1]")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp)
        path = Path(tmp.name)

    try:
        result = await handler.handle_deanonymize_file("s1", str(path))

        # Перечитываем собранный DOCX
        d2 = Document(str(path))
        text = "\n".join(p.text for p in d2.paragraphs)
        assert "[PERSON_1]" not in text, text
        assert "Ивана Петрова" in text, text
        print("TEST 2 OK: deanonymize_file .docx — плейсхолдер заменён, структура сохранена")
    finally:
        path.unlink(missing_ok=True)


async def test_deanonymize_file_missing():
    store = FakeStore()
    handler = make_handler(store)
    try:
        await handler.handle_deanonymize_file("s1", "c:/no/such/file.md")
        raise AssertionError("ожидался FileNotFoundError")
    except FileNotFoundError:
        pass
    print("TEST 3 OK: deanonymize_file — отсутствующий файл -> FileNotFoundError")


async def main():
    await test_deanonymize_file_md()
    await test_deanonymize_file_docx()
    await test_deanonymize_file_missing()
    print("\nALL DEANONYMIZE FILE TESTS PASSED")


if __name__ == "__main__":
    asyncio.run(main())
