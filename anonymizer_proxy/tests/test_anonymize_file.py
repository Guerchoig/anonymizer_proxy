"""
Функциональные тесты анонимизации локального файла (/api/anonymize_file).

Проверяют создание анонимизированной копии файла (с плейсхолдерами) рядом
с оригиналом, с сохранением структуры (.docx).

Запуск: python anonymizer_proxy\\tests\\test_anonymize_file.py (из корня проекта)
"""
import asyncio
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from anonymizer_proxy.proxy.handlers import RequestHandler
from anonymizer_proxy.tests.test_tools_manual import FakeNER, FakeStore, FakeOpenRouter


async def test_anonymize_file_docx():
    from docx import Document

    store = FakeStore()
    handler = RequestHandler(
        ner_service=FakeNER({"Ивана Петрова": "PERSON"}),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )

    doc = Document()
    doc.add_paragraph("Документ подготовлен от Ивана Петрова")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ФИО"
    table.cell(0, 1).text = "Должность"
    table.cell(1, 0).text = "Ивана Петрова"
    table.cell(1, 1).text = "Директор"
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp)
        path = Path(tmp.name)

    anon_path = None
    try:
        result = await handler.handle_anonymize_file(str(path))
        anon_path = Path(result["anonymized_file"])

        assert anon_path.exists(), result
        assert anon_path.name.endswith(".anonymized.docx"), anon_path.name
        assert anon_path.parent == path.parent

        # Копия содержит плейсхолдер, а не оригинальное имя
        d2 = Document(str(anon_path))
        anon_text = "\n".join(p.text for p in d2.paragraphs)
        assert "[PERSON_1]" in anon_text, anon_text
        assert "Ивана Петрова" not in anon_text, anon_text

        # Оригинал не изменён
        d3 = Document(str(path))
        orig_text = "\n".join(p.text for p in d3.paragraphs)
        assert "Ивана Петрова" in orig_text, orig_text

        # Markdown-представление с таблицами: плейсхолдер в markdown-таблице
        md = result["anonymized_markdown"]
        assert "[PERSON_1]" in md, md
        assert "Ивана Петрова" not in md, md
        assert "| ---" in md, md  # markdown-таблица
        assert result["review_file"] is not None

        print("TEST 1 OK: anonymize_file .docx — копия с плейсхолдером, markdown с таблицами")
    finally:
        path.unlink(missing_ok=True)
        if anon_path:
            anon_path.unlink(missing_ok=True)


async def test_anonymize_file_output_path():
    from docx import Document

    store = FakeStore()
    handler = RequestHandler(
        ner_service=FakeNER({"Ивана Петрова": "PERSON"}),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )
    doc = Document()
    doc.add_paragraph("Текст Ивана Петрова")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp)
        path = Path(tmp.name)

    out = path.with_name("custom_output.docx")
    try:
        result = await handler.handle_anonymize_file(str(path), output_path=str(out))
        assert Path(result["anonymized_file"]) == out
        assert out.exists()
        print("TEST 2 OK: anonymize_file — явный output_path")
    finally:
        path.unlink(missing_ok=True)
        out.unlink(missing_ok=True)


async def test_anonymize_file_missing():
    store = FakeStore()
    handler = RequestHandler(
        ner_service=FakeNER({}),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )
    try:
        await handler.handle_anonymize_file("c:/no/such.docx")
        raise AssertionError("ожидался FileNotFoundError")
    except FileNotFoundError:
        pass
    print("TEST 3 OK: anonymize_file — отсутствующий файл -> FileNotFoundError")


async def test_cross_segment_entity_keeps_alignment():
    """Регрессия 2026-09-08: NER-сущность, пересекающая границу сегментов
    склеенного текста (в реальном багрепорте «Sasha\\nА.В. Гершойг» — автор
    и текст комментария; «АП\\nОбособленные подразделения» — многоабзацная
    ячейка), раньше заменялась в склеенном тексте целиком: «\\n» внутри
    спана удалялся, число строк анонимизированного текста переставало
    совпадать с числом сегментов, и хвост сегментов (вложенные таблицы,
    комментарии) оставался без замен. Теперь замена выполняется
    посегментно — соответствие «строка ↔ сегмент» сохраняется."""
    from docx import Document

    store = FakeStore()
    handler = RequestHandler(
        ner_service=FakeNER({
            "АП\nОбособленные подразделения": "DEPARTMENT",
            "Ивана Петрова": "PERSON",
        }),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )

    doc = Document()
    doc.add_paragraph("Начало документа")
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].text = "АП"
    cell.add_paragraph("Обособленные подразделения")
    doc.add_paragraph("Подписал Ивана Петрова")  # хвост — маркер выравнивания
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp)
        path = Path(tmp.name)

    anon_path = None
    try:
        result = await handler.handle_anonymize_file(str(path))
        anon_path = Path(result["anonymized_file"])

        d2 = Document(str(anon_path))
        # Хвостовой сегмент (после кросс-сегментной замены) анонимизирован —
        # выравнивание «строка ↔ сегмент» не сломано
        tail = d2.paragraphs[-1].text
        assert "[PERSON_" in tail, tail
        assert "Ивана Петрова" not in tail, tail
        # Обе части кросс-сегментной сущности заменены своими токенами
        cell_text = "\n".join(
            p.text for p in d2.tables[0].cell(0, 0).paragraphs)
        assert "АП" not in cell_text, cell_text
        assert "Обособленные подразделения" not in cell_text, cell_text
        assert "[DEPARTMENT_" in cell_text, cell_text
        # Обычный абзац не пострадал
        assert d2.paragraphs[0].text == "Начало документа"
        print("TEST 4 OK: кросс-сегментная сущность — выравнивание строк "
              "сохранено, все вхождения заменены")
    finally:
        path.unlink(missing_ok=True)
        if anon_path:
            anon_path.unlink(missing_ok=True)


async def test_stale_copy_invalidated_by_marker():
    """Регрессия 2026-09-08: копия, созданная СТАРОЙ версией анонимизатора
    (без маркера версии), должна пересоздаваться при повторной команде, а не
    переиспользоваться — иначе копии, сделанные до обновления парсера,
    навсегда остаются неполными (утечки во вложенных таблицах/комментариях)."""
    from docx import Document
    from anonymizer_proxy.anonymizer.file_parser import ANON_MARKER

    store = FakeStore()
    handler = RequestHandler(
        ner_service=FakeNER({"Ивана Петрова": "PERSON"}),
        mapping_store=store,
        openrouter_client=FakeOpenRouter(),
    )
    doc = Document()
    doc.add_paragraph("Документ от Ивана Петрова")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp)
        path = Path(tmp.name)

    anon_path = None
    try:
        r1 = await handler.handle_anonymize_file(str(path))
        anon_path = Path(r1["anonymized_file"])
        assert not r1.get("reused", False)
        # Копия несёт маркер версии
        assert Document(str(anon_path)).core_properties.comments == ANON_MARKER

        # Повторная команда по неизменённому исходнику — переиспользование
        r2 = await handler.handle_anonymize_file(str(path))
        assert r2.get("reused") is True, r2

        # Имитируем копию от старой версии: стираем маркер
        d = Document(str(anon_path))
        d.core_properties.comments = ""
        buf = io.BytesIO()
        d.save(buf)
        anon_path.write_bytes(buf.getvalue())

        # Третья команда — копия ПЕРЕСОЗДАЁТСЯ (полный прогон, не reuse)
        r3 = await handler.handle_anonymize_file(str(path))
        assert not r3.get("reused", False), r3
        assert Path(r3["anonymized_file"]).exists()
        assert Document(str(anon_path)).core_properties.comments == ANON_MARKER
        print("TEST 5 OK: копия без маркера версии пересоздаётся, "
              "с маркером — переиспользуется")
    finally:
        path.unlink(missing_ok=True)
        if anon_path:
            anon_path.unlink(missing_ok=True)


async def main():
    await test_anonymize_file_docx()
    await test_anonymize_file_output_path()
    await test_anonymize_file_missing()
    await test_cross_segment_entity_keeps_alignment()
    await test_stale_copy_invalidated_by_marker()
    print("\nALL ANONYMIZE FILE TESTS PASSED")


if __name__ == "__main__":
    asyncio.run(main())
